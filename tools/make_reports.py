"""Build ЕСЕП.md and ЕСЕП.docx for each lab from its script, output and metadata."""

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from labs_meta import COMMON_TOOLS, LABS  # noqa: E402
from labs_qa import ANSWERS, EXTRA_LIT, LITERATURE  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent

# Автор/оқытушы деректері репозиторийде сақталмайды: tools/student.local.json (git-ке кірмейді) ішінен оқылады.
# Үлгі: tools/student.example.json. Файл жоқ болса — плейсхолдерлер қойылады.
IDENTITY = {
    "course": "Үлкен деректерді талдау",
    "program": "<білім беру бағдарламасы, курс, топ>",
    "teacher": "<оқытушының аты-жөні>",
    "student": "<студенттің аты-жөні>",
}
_id_file = Path(__file__).parent / "student.local.json"
if _id_file.exists():
    IDENTITY.update(json.loads(_id_file.read_text(encoding="utf-8")))
DATE = "2026-10-05"
MARK_SCRIPT = re.compile(r"^# ==== (.+?) ====\s*$")
MARK_OUT = re.compile(r"^> # ==== (.+?) ====\s*$")
NOISE = re.compile(r"^(Attaching package|The following objects? (is|are) masked|    \S+, |    %notin%|    margin|    combine|randomForest 4|Type rfNews)")
MAX_OUT = 42


def task_number(title: str):
    m = re.match(r"(\d+)-", title)
    return int(m.group(1)) if m else None


def script_sections(text: str) -> list[tuple[str, str]]:
    sections, title, buf = [], None, []
    for line in text.splitlines():
        m = MARK_SCRIPT.match(line)
        if m:
            if title:
                sections.append((title, "\n".join(buf).strip()))
            title, buf = m.group(1), []
        elif title:
            buf.append(line)
    if title:
        sections.append((title, "\n".join(buf).strip()))
    return sections


def script_preamble(text: str) -> str:
    lines = []
    for line in text.splitlines():
        if MARK_SCRIPT.match(line):
            break
        lines.append(line)
    return "\n".join(lines).strip()


def output_sections(text: str) -> dict[str, list[str]]:
    res, title = {}, None
    for line in text.splitlines():
        m = MARK_OUT.match(line)
        if m:
            title = m.group(1)
            res[title] = []
            continue
        if title is None or line.startswith(("> ", "+ ")) or line in (">", "+") or NOISE.match(line):
            continue
        res[title].append(line.replace("\t", "    "))
    for k, v in res.items():
        while v and not v[-1].strip():
            v.pop()
        while v and not v[0].strip():
            v.pop(0)
        if len(v) > MAX_OUT:
            res[k] = v[:MAX_OUT] + [f"... (тағы {len(v) - MAX_OUT} жол, толық нәтиже output.txt файлында)"]
    return res


def build(lab_dir: Path) -> None:
    meta = LABS[lab_dir.name]
    n = meta["n"]
    script_path = next(lab_dir.glob("lab*.R"))
    script = script_path.read_text(encoding="utf-8")
    outs = output_sections((lab_dir / "output.txt").read_text(encoding="utf-8"))
    sections = script_sections(script)
    shots = sorted((lab_dir / "screenshots").glob("console_*.png"))
    questions = meta["questions"]
    answers = ANSWERS[n]
    literature = LITERATURE + EXTRA_LIT.get(n, [])

    # ---------- Markdown ----------
    md = [
        f"# Зертханалық сабақ №{n}", "",
        f"**Тақырыбы:** {meta['title']}", "",
        f"**Пәні:** {IDENTITY['course']}  ",
        f"**ББ:** {IDENTITY['program']}  ",
        f"**Оқытушы:** {IDENTITY['teacher']}  ",
        f"**Орындаушы:** {IDENTITY['student']}  ",
        f"**Күні:** {DATE}", "", "---", "",
        "## 1. Жұмыстың мақсаты", "", meta["goal"], "",
        "## 2. Теориялық бөлім", "",
    ]
    for p in meta["theory"]:
        md += [p, ""]
    md += ["## 3. Қолданылған құралдар", "",
           f"- **Орта:** {COMMON_TOOLS}",
           f"- **Пакеттер:** {meta['tools']}",
           f"- **Скрипт:** `{script_path.name}` (толық нәтиже — `output.txt`)", ""]
    pre = script_preamble(script)
    md += ["## 4. Жұмыстың орындалу барысы", "", "Скрипттің басы (пакеттерді қосу):", "", "```r", pre, "```", ""]
    for i, (title, code) in enumerate(sections, 1):
        md += [f"### 4.{i}. {title}", "", "```r", code, "```", ""]
        out = outs.get(title, [])
        if out:
            md += ["Нәтиже:", "", "```text", *out, "```", ""]
        c = meta["comments"].get(task_number(title))
        if c:
            md += [f"**Түсіндірме:** {c}", ""]
    if "extra_table" in meta:
        cap, head, rows = meta["extra_table"]
        md += [f"**{cap}:**", "", "| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
        md += ["| " + " | ".join(r) + " |" for r in rows]
        md += [""]
    md += ["## 5. Скриншоттар", "", "### 5.1. RStudio консолі", ""]
    for s in shots:
        md += [f"![{s.stem}](screenshots/{s.name})", ""]
    if meta["figures"]:
        md += ["### 5.2. Графиктер", ""]
        for i, (f, cap) in enumerate(meta["figures"], 1):
            md += [f"**{i}-сурет.** {cap}", "", f"![{cap}](screenshots/{f})", ""]
    md += ["## 6. Бақылау сұрақтарына жауаптар", ""]
    for i, (q, a) in enumerate(zip(questions, answers), 1):
        md += [f"**{i}. {q}**", "", a, ""]
    md += ["## 7. Қорытынды", "", meta["conclusion"], "",
           "## 8. Қолданылған әдебиеттер", ""]
    md += [f"{i}. {lit}" for i, lit in enumerate(literature, 1)]
    md += [""]
    (lab_dir / "ЕСЕП.md").write_text("\n".join(md), encoding="utf-8")

    # ---------- Word ----------
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(4)
    for hn, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
        h = doc.styles[hn]
        h.font.name, h.font.size, h.font.bold = "Times New Roman", Pt(size), True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def para(text, bold_prefix=None, align=None, size=None, italic=False):
        p = doc.add_paragraph()
        if align:
            p.alignment = align
        parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
        if bold_prefix:
            r = p.add_run(bold_prefix)
            r.bold = True
        for part in parts:
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(12)
            else:
                r = p.add_run(part)
            if size:
                r.font.size = Pt(size)
            r.italic = italic
        return p

    def code(text, label=None):
        if label:
            para(label, size=12, italic=True)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name, r.font.size = "Courier New", Pt(9)
        r.element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        shd = p._p.get_or_add_pPr().makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "F2F2F2"})
        p._p.get_or_add_pPr().append(shd)

    t = doc.add_heading(f"Зертханалық сабақ №{n}", level=1)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(meta["title"], bold_prefix="Тақырыбы: ")
    para(IDENTITY["course"], bold_prefix="Пәні: ")
    para(IDENTITY["program"], bold_prefix="ББ: ")
    para(IDENTITY["teacher"], bold_prefix="Оқытушы: ")
    para(IDENTITY["student"], bold_prefix="Орындаушы: ")
    para(DATE, bold_prefix="Күні: ")
    doc.add_heading("1. Жұмыстың мақсаты", level=2)
    para(meta["goal"])
    doc.add_heading("2. Теориялық бөлім", level=2)
    for p in meta["theory"]:
        para(p)
    doc.add_heading("3. Қолданылған құралдар", level=2)
    para(COMMON_TOOLS, bold_prefix="Орта: ")
    para(meta["tools"], bold_prefix="Пакеттер: ")
    para(f"`{script_path.name}` (толық нәтиже — `output.txt`)", bold_prefix="Скрипт: ")
    doc.add_heading("4. Жұмыстың орындалу барысы", level=2)
    code(pre, "Скрипттің басы (пакеттерді қосу):")
    for i, (title, c_) in enumerate(sections, 1):
        doc.add_heading(f"4.{i}. {title}", level=3)
        code(c_, "Код:")
        out = outs.get(title, [])
        if out:
            code("\n".join(out), "Нәтиже:")
        c = meta["comments"].get(task_number(title))
        if c:
            para(c, bold_prefix="Түсіндірме: ")
    if "extra_table" in meta:
        cap, head, rows = meta["extra_table"]
        para(cap + ":", size=13).runs[0].bold = True
        tb = doc.add_table(rows=1, cols=len(head))
        tb.style = "Table Grid"
        for j, h in enumerate(head):
            tb.rows[0].cells[j].text = h
            tb.rows[0].cells[j].paragraphs[0].runs[0].bold = True
        for r in rows:
            cells = tb.add_row().cells
            for j, v in enumerate(r):
                cells[j].text = v
    doc.add_heading("5. Скриншоттар", level=2)
    doc.add_heading("5.1. RStudio консолі", level=3)
    for s in shots:
        doc.add_picture(str(s), width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if meta["figures"]:
        doc.add_heading("5.2. Графиктер", level=3)
        for i, (f, cap) in enumerate(meta["figures"], 1):
            doc.add_picture(str(lab_dir / "screenshots" / f), width=Cm(15))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para(cap, bold_prefix=f"{i}-сурет. ", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    doc.add_heading("6. Бақылау сұрақтарына жауаптар", level=2)
    for i, (q, a) in enumerate(zip(questions, answers), 1):
        para(f"**{i}. {q}**")
        para(a)
    doc.add_heading("7. Қорытынды", level=2)
    for p in meta["conclusion"].split("\n\n"):
        para(p)
    doc.add_heading("8. Қолданылған әдебиеттер", level=2)
    for i, lit in enumerate(literature, 1):
        para(f"{i}. {lit}")
    doc.save(lab_dir / "ЕСЕП.docx")
    print(f"{lab_dir.name}: {len(sections)} тапсырма, {len(shots)} консоль, {len(meta['figures'])} график")


if __name__ == "__main__":
    targets = [ROOT / a.rstrip("/") for a in sys.argv[1:]] or sorted(ROOT.glob("lab*/"))
    for d in targets:
        build(d)
